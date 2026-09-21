"""
Fleet Speeding Alert Summary — weekly tool
Upload the raw Vehicle Exception Report; explore an overview + per-driver drill-down; download.

Run locally:
    pip install streamlit pandas openpyxl altair
    streamlit run speeding_app.py
"""
import io
import pandas as pd
import streamlit as st
import altair as alt
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ---- speeding rules: flag when Speed >= threshold for that limit ----
THRESHOLDS = {20: 35, 30: 44, 40: 56, 50: 66, 60: 76, 70: 91}

TEAL = "0E6E63"; WHITE = "FFFFFF"; BORD = "D5DEDC"
thin = Side(style="thin", color=BORD); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
ACCENT = "#0E6E63"


def load_raw(uploaded) -> pd.DataFrame:
    name = uploaded.name.lower()
    df = pd.read_excel(uploaded, dtype=str) if name.endswith((".xlsx", ".xls")) else pd.read_csv(uploaded, dtype=str)
    df.columns = [str(c).strip() for c in df.columns]
    return df


def build_flagged(df: pd.DataFrame, window_min: int = 15):
    """Return the full flagged-event detail (with alert grouping) + meta."""
    need = ["Driver", "Vehicle", "Date", "Time", "Speed", "Speed Limit"]
    missing = [c for c in need if c not in df.columns]
    if missing:
        raise ValueError(f"Missing expected column(s): {', '.join(missing)}")

    d = df.copy()
    for c in ["Driver", "Vehicle"]:
        d[c] = d[c].astype(str).str.strip()
    d["Speed"] = pd.to_numeric(d["Speed"], errors="coerce")
    d["Speed Limit"] = pd.to_numeric(d["Speed Limit"], errors="coerce")
    d["dt"] = pd.to_datetime(d["Date"].astype(str).str.strip() + " " + d["Time"].astype(str).str.strip(),
                             dayfirst=True, errors="coerce")
    d = d[d["dt"].notna() & d["Speed"].notna() & d["Speed Limit"].notna() & (d["Driver"] != "")]

    d["_thr"] = d["Speed Limit"].map(THRESHOLDS)
    d = d[d["_thr"].notna() & (d["Speed"] >= d["_thr"])].sort_values(["Driver", "dt"]).copy()

    # continuous 15-min rule per driver -> new alert when gap > window
    gap = d.groupby("Driver")["dt"].diff().dt.total_seconds()
    d["New alert"] = ((gap.isna()) | (gap > window_min * 60)).astype(int)
    d["Alert #"] = d.groupby("Driver")["New alert"].cumsum()
    d["day"] = d["dt"].dt.date
    d["Over limit (mph)"] = (d["Speed"] - d["Speed Limit"]).astype(int)

    meta = {"date_from": d["day"].min(), "date_to": d["day"].max()}
    return d, meta


def build_summary(flagged: pd.DataFrame) -> pd.DataFrame:
    """Driver summary (with serial no + per-day real alerts) from a (filtered) flagged set."""
    if flagged.empty:
        return pd.DataFrame()
    days = sorted(flagged["day"].unique())
    day_cols = {dd: dd.strftime("%a %d/%m") for dd in days}
    rows = []
    for drv, g in flagged.groupby("Driver"):
        veh = g["Vehicle"].mode()
        row = {"Driver": drv, "Vehicle Reg": veh.iloc[0] if len(veh) else "",
               "Speeding events": int(len(g)), "Real Alerts": int(g["New alert"].sum())}
        per_day = g.groupby("day")["New alert"].sum()
        for dd in days:
            row[day_cols[dd]] = int(per_day.get(dd, 0))
        rows.append(row)
    summ = pd.DataFrame(rows).sort_values("Speeding events", ascending=False).reset_index(drop=True)
    summ.insert(0, "S/No", range(1, len(summ) + 1))
    return summ


def to_excel(summary: pd.DataFrame, detail: pd.DataFrame) -> bytes:
    buf = io.BytesIO()
    det = detail[["Vehicle", "Driver", "Date", "Time", "Speed", "Speed Limit", "Over limit (mph)", "Alert #", "New alert"]].copy()
    with pd.ExcelWriter(buf, engine="openpyxl") as xl:
        summary.to_excel(xl, sheet_name="Driver Summary", index=False)
        det.to_excel(xl, sheet_name="Flagged detail", index=False)
        pd.DataFrame({"Speed limit": list(THRESHOLDS), "Flag at (mph)": list(THRESHOLDS.values())}
                     ).to_excel(xl, sheet_name="Rules", index=False)
        for sn in xl.book.sheetnames:
            ws = xl.book[sn]
            for cell in ws[1]:
                cell.font = Font("Calibri", bold=True, color=WHITE, size=10)
                cell.fill = PatternFill("solid", fgColor=TEAL)
                cell.alignment = Alignment("center", "center", wrap_text=True); cell.border = BORDER
            ws.row_dimensions[1].height = 26
            for col in ws.columns:
                w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[col[0].column_letter].width = min(max(w + 2, 10), 30)
            ws.freeze_panes = "A2"
    return buf.getvalue()


# ------------------------------- UI -------------------------------
st.set_page_config(page_title="Fleet Speeding Alerts", page_icon="🚗", layout="wide")
st.title("🚗 Fleet speeding alert summary")
st.caption("Upload this week's Vehicle Exception Report, explore the overview, and drill into any driver.")

with st.sidebar:
    st.subheader("Rules")
    st.markdown("**Flag speed at:** 20→35 · 30→44 · 40→56 · 50→66 · 60→76 · 70→91 mph")
    window_min = st.number_input("Group events within (minutes) = 1 alert", 1, 60, 15, 1)

up = st.file_uploader("Drop your raw report here (CSV or Excel)", type=["csv", "xlsx", "xls"])
if up is None:
    st.info("Waiting for a file… drag your weekly report above.")
    st.stop()

try:
    flagged_all, meta = build_flagged(load_raw(up), window_min)
except Exception as e:
    st.error(f"Could not process the file: {e}")
    st.stop()

if flagged_all.empty:
    st.warning("No speeding events met the flag thresholds in this file.")
    st.stop()

# ---------------- Filters ----------------
days_all = sorted(flagged_all["day"].unique())
st.markdown("#### Filters")
f1, f2, f3 = st.columns([1.2, 1, 1])
with f1:
    dr = st.date_input("Date range", (days_all[0], days_all[-1]),
                       min_value=days_all[0], max_value=days_all[-1])
    if isinstance(dr, (list, tuple)) and len(dr) == 2:
        d_from, d_to = dr
    else:
        d_from = d_to = dr if not isinstance(dr, (list, tuple)) else dr[0]
with f2:
    regs = st.multiselect("Vehicle reg", sorted(flagged_all["Vehicle"].unique()))
with f3:
    drvs = st.multiselect("Driver", sorted(flagged_all["Driver"].unique()))

fl = flagged_all[(flagged_all["day"] >= d_from) & (flagged_all["day"] <= d_to)]
if regs:
    fl = fl[fl["Vehicle"].isin(regs)]
if drvs:
    fl = fl[fl["Driver"].isin(drvs)]

if fl.empty:
    st.warning("No events match the current filters.")
    st.stop()

tab_overview, tab_drill = st.tabs(["📊 Overview", "🔎 Driver drill-down"])

# ---------------- Overview ----------------
with tab_overview:
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Drivers", fl["Driver"].nunique())
    c2.metric("Flagged events", f"{len(fl):,}")
    c3.metric("Real alerts", f"{int(fl['New alert'].sum()):,}")
    c4.metric("Window", f"{d_from.strftime('%d/%m')} – {d_to.strftime('%d/%m')}")

    summary = build_summary(fl)

    st.subheader("Top 10 offenders")
    rank_by = st.radio("Rank by", ["Real Alerts", "Speeding events"], horizontal=True)
    top = summary.nlargest(10, rank_by)[["Driver", rank_by]]
    ch = (alt.Chart(top).mark_bar(color=ACCENT, cornerRadiusEnd=4)
          .encode(x=alt.X(f"{rank_by}:Q", title=rank_by),
                  y=alt.Y("Driver:N", sort="-x", title=None), tooltip=["Driver", rank_by])
          .properties(height=340))
    st.altair_chart(ch + ch.mark_text(align="left", dx=4, color="#0A4A43").encode(text=f"{rank_by}:Q"),
                    use_container_width=True)

    st.subheader("Driver Summary")
    st.dataframe(summary, use_container_width=True, hide_index=True)
    st.download_button("⬇️ Download summary (Excel)", to_excel(summary, fl),
                       file_name=f"Driver_Speeding_Summary_{d_from:%d-%m}_{d_to:%d-%m}.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

# ---------------- Driver drill-down ----------------
with tab_drill:
    order = fl.groupby("Driver")["New alert"].sum().sort_values(ascending=False).index.tolist()
    who = st.selectbox("Choose a driver (sorted by real alerts)", order)
    g = fl[fl["Driver"] == who].sort_values("dt")

    k1, k2, k3, k4 = st.columns(4)
    k1.metric("Speeding events", len(g))
    k2.metric("Real alerts", int(g["New alert"].sum()))
    k3.metric("Worst over limit", f"{int(g['Over limit (mph)'].max())} mph")
    k4.metric("Avg over limit", f"{g['Over limit (mph)'].mean():.0f} mph")

    a, b = st.columns(2)
    with a:
        st.markdown("**Real alerts by day**")
        byday = g[g["New alert"] == 1].groupby("day").size().reset_index(name="Alerts")
        byday["day"] = byday["day"].astype(str)
        st.altair_chart(alt.Chart(byday).mark_bar(color=ACCENT).encode(
            x=alt.X("day:N", title=None), y=alt.Y("Alerts:Q")).properties(height=260),
            use_container_width=True)
    with b:
        st.markdown("**Events by road speed limit**")
        bylim = g.groupby("Speed Limit").size().reset_index(name="Events")
        st.altair_chart(alt.Chart(bylim).mark_bar(color="#1F7A8C").encode(
            x=alt.X("Speed Limit:O", title="Speed limit (mph)"), y=alt.Y("Events:Q")).properties(height=260),
            use_container_width=True)

    st.markdown("**How far over the limit (mph)**")
    st.altair_chart(alt.Chart(g).mark_bar(color="#E4572E").encode(
        x=alt.X("Over limit (mph):Q", bin=alt.Bin(step=5), title="mph over the limit"),
        y=alt.Y("count():Q", title="Events")).properties(height=240), use_container_width=True)

    st.subheader(f"Raw speeding events — {who}")
    det = g[["Date", "Time", "Vehicle", "Speed", "Speed Limit", "Over limit (mph)", "Alert #"]].reset_index(drop=True)
    det.insert(0, "S/No", range(1, len(det) + 1))
    st.dataframe(det, use_container_width=True, hide_index=True)
    st.download_button("⬇️ Download this driver's events (CSV)", det.to_csv(index=False).encode(),
                       file_name=f"{who.replace(' ','_')}_events.csv", mime="text/csv")
